"""Build the report from the current reproducible synthetic evaluation."""
from __future__ import annotations

import csv
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from evaluate_report import build_evaluation

INPUT = ROOT / "data/processed/pedestrian_crossings_fuzzy_topsis.csv"
OUT = ROOT / "reports/evaluation"
FIGS = ROOT / "reports/report_figures"
evaluation = build_evaluation(INPUT, OUT, draws=1000)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Times New Roman", Pt(10.5)
normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(5), 1.03
for name, size in (("Title", 19), ("Heading 1", 13), ("Heading 2", 11)):
    doc.styles[name].font.name = "Times New Roman"
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)

def p(text):
    doc.add_paragraph(str(text))

def h(text, level=1):
    doc.add_heading(text, level=level)

def table(headers, rows, size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    for c, value in zip(t.rows[0].cells, headers): c.text = str(value)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, row): c.text = str(value)
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs: run.font.size = Pt(size)
    return t

def f(value, digits=4): return f"{float(value):.{digits}f}"

title = "A Fuzzy TOPSIS Framework for Prioritising Pedestrian Crossing Safety and Accessibility Improvements"
doc.add_paragraph(title, "Title")
p("Team [Team number] • [Team member names]")
p("Department of [Department], IIIT Kottayam, India • [Institutional email addresses]")
h("Abstract")
p(f"This report asks whether adding essential-destination access (E1) changes the leading group beyond a four-criterion baseline. In the 28-candidate simulation, the five-criterion model gives a tied top score of {f(evaluation['all_candidates'][0]['score'],4)}; the top-three overlap with the four-criterion model is {f(evaluation['comparisons'][1]['top_k_overlap'],3)}, and 1,000 bounded weight draws have mean top-three overlap {f(evaluation['sensitivity']['mean_top_k_overlap'],3)}. The pedestrian-demand/E1 central-score Spearman correlation is {f(evaluation['ped_destination_spearman'],2)}, indicating substantial constructed-proxy overlap. These are synthetic calculations informed by public map geometry—not measured conditions, verified crossings, or an empirical safety ranking. The older four-row example remains a separate teaching illustration.")
p("Keywords: pedestrian crossings; accessibility; fuzzy TOPSIS; synthetic data; sensitivity analysis")

h("I Introduction")
p("A transparent multi-criteria ranking can organise review when demand, road context, accessibility and control evidence are incomplete, but its meaning depends on its inputs and candidate set. This project ranks hypothetical improvement packages for earlier review—not danger, crash prevention, or a budget-optimal portfolio. Its contribution is a reproducible context-informed synthetic dataset: unlike independently random alternatives, candidate selection follows saved road geometry and shared-node junction topology, while road hierarchy and mapped destinations inform explicit proxy formulas. This creates traceable intermediate evidence and assumptions, not stronger empirical observations; modeled deficiencies remain prior-based and every score remains synthetic. Feasibility is an external engineering and jurisdictional gate.")
h("II Related Work")
p("Structured crossing checklists [1], pedestrian-facility prioritisation [2], and fuzzy gap-acceptance analysis [3] motivate documented evidence and decision support, but their settings and outcomes do not validate these inputs. Chen [4] informs the TOPSIS lineage; fuzzy conversion research [5] motivates explicit sensitivity. Kerala’s road-safety vision [6] and a midblock study from four other Indian cities [7] provide context only. OSM [8] supplies the mapped layer; Kerala district charts [9] are not score inputs. The declared calculations below are project adaptations, not an unchanged reproduction of an earlier formulation.")
table(["Study", "Data and approach", "Relevance and transfer limit"], [
    ["Battistini et al. [1]", "Checklists at ten unsignalized crossings on two Bologna roads; decision support for intervention priorities.", "Supports structured inspection; Italian site ratings are not local observations or Indian compliance thresholds."],
    ["Majumdar et al. [2]", "Pedestrian satisfaction and TOPSIS prioritisation of sidewalk/crosswalk attributes in an Indian context.", "Supports pedestrian-focused criteria; ranking infrastructure attributes differs from ranking crossing packages."],
    ["Alver and Katanalp [3]", "Integrated AHP and fuzzy logic for gap acceptance at semicontrolled midblock crossings.", "Illustrates fuzzy transport modelling; behavioral analysis does not validate our synthetic deficiency scale."],
], size=9)

doc.add_page_break()
h("III Study Scope and Data")
p("The generator reads the saved data/raw/osm_context.json extract and data/raw/osm_metadata.json metadata. The extract contains 431 elements (286 accepted motor-road ways; 91 deduplicated destinations). From the mapped center (9.7554364° N, 76.6495350° E), a deterministic 3 km search, ≥80 m spacing, and map-evidence priority selected 28 coordinates: 27 road-junction candidates and one explicitly mapped crossing. All are unverified. Shared-node junctions are candidate locations, not proof of crossings; road ownership, field conditions, and map completeness are unknown.")
p("Evidence chain: saved map geometry/tags and destination objects are mapped facts; formulas derive context indicators; deterministic seeded rules produce synthetic ordinal centers and ranges. Example SYN-OSM-716872373 is near mapped St Mary’s Church (60 m), Alphonsa English Medium School Valavoor (83 m), Valavoor School (277 m), and a post office (286 m), all straight-line distances. Proximity does not show a walking route or access need. All centers are integer scores 1–5. Fuzzy widths are ±1.0 for low-confidence proxies and ±1.5 for very-low/unknown assumptions, clipped to scale bounds; these are not confidence intervals.")
table(["Criterion", "Map/context input and generation rule", "Interpretation"], [
    ["P1 Demand", "Nearby mapped amenity intensity I, bus stop B, vehicle score; round_clip(1+1.05 ln(1+I)+0.35B+0.12(vehicle−2)+U(−0.65,0.65)).", "Constructed demand proxy; no counts."],
    ["P2 Vehicle", "Highest adjacent road class (motorway/trunk/primary 5; secondary 4; tertiary 3; residential/unclassified 2; service/living street 1); +0.5 for each satisfied condition: ≥3 mapped lanes, numeric maxspeed ≥60 km/h; U(−0.35,0.35).", "Road context, not volume or operating speed."],
    ["P3 Access", "Prior 3 + shared U(−0.45,0.45) + N(0,0.85); mapped wheelchair, kerb/sloped curb and tactile tags adjust it.", "Unknown tags stay unknown; deficiency mostly assumed."],
    ["P4 Control", "Prior 3 + shared variation +0.2(vehicle−3)+N(0,0.85); mapped signals/markings and lighting adjust it.", "Not a visibility or condition audit."],
    ["E1 Destination", "Strongest mapped destination proximity plus total intensity: round_clip(1+1.05M+0.25 ln(1+I)).", "Straight-line proxy, not walking-network access."],
], size=8.2)
p("Within 500 m, each destination contributes a(1−d/600), where d is great-circle distance in metres and a is the declared destination-type weight. I sums these contributions; M is their maximum (zero when none are mapped). B=1 if a mapped bus stop/platform lies within 250 m, otherwise 0. Schools/clinics/hospitals have a=3, universities/pharmacies/markets 2.5, bus stops 2, banks/post offices/community or worship facilities 1.5, and other shops 1. round_clip rounds to the nearest integer and clips to 1–5. N(0,0.85) denotes mean zero and standard deviation 0.85; all draws are seeded assumptions.")
p("Scores/ranges are synthetic, not calibrated. For deficiency criteria, the prior is 3 and missing tags do not imply absence. Generator seed 20261024 differs from sensitivity seed 20261009. Exact feature weights, road-class constants, draw order, and source records are in reports/synthetic_data_methodology.md and data/source_inventory.csv. No counts, surveys, operating speeds, expert validation, or feasibility findings were invented.")

doc.add_page_break()
h("IV Methodology")
p("For alternative i and criterion j, represent the priority-oriented input as a triangular fuzzy number x̃ᵢⱼ=(l,c,u), where l≤c≤u. Components are normalized against shared input outer bounds Lⱼ=minᵢlᵢⱼ and Uⱼ=maxᵢuᵢⱼ: z=(x−Lⱼ)/(Uⱼ−Lⱼ). Non-discriminating criteria are removed and remaining equal weights renormalized. With five active criteria, each weight is 1/5. Weighted triangles are compared with theoretical positive ideals (wⱼ,wⱼ,wⱼ) and negative ideals (0,0,0). The positive ideal means greatest intervention priority, not the safest condition.")
p("The vertex distance is d(ã,b̃)=√(((aₗ−bₗ)²+(a꜀−b꜀)²+(aᵤ−bᵤ)²)/3). Criterion distances are summed: Dᵢ⁺=Σⱼd(ṽᵢⱼ,Aⱼ⁺), Dᵢ⁻=Σⱼd(ṽᵢⱼ,Aⱼ⁻); closeness is Cᵢ=Dᵢ⁻/(Dᵢ⁺+Dᵢ⁻). Descending C ranks higher priority. This is a declared project adaptation: shared-range normalization, equal weights, weighted theoretical ideals, and summed distances. Fuzzy membership grades are not probabilities; ranges are uncalibrated stress-test representations.")
p(f"Research question: does adding E1 change ordering while the four original criteria are retained? E1 is a 500 m straight-line proxy, not the field definition based on verified walking routes and access necessity. Demand and E1 share amenity context; their central-score Spearman correlation is {f(evaluation['ped_destination_spearman'],3)}, so double-counting is plausible. The crisp comparator min–max scales central values and averages active columns; it is not crisp TOPSIS. The comparison does not establish predictive accuracy.")
h("Processing path and comparison design", 2)
table(["Stage", "Input → output", "What enters model"], [
    ["Select candidates", "osm_context.json + osm_metadata.json → mapped candidates/context", "Coordinates, road/destination context"],
    ["Generate synthetic rows", "build_synthetic_dataset.py → processed/pedestrian_crossings_fuzzy_topsis.csv", "Five synthetic criterion triples"],
    ["Run CLI", "--synthetic-input reads that CSV → scores/ranks", "Criterion triples only; metadata fields validate provenance, not criteria"],
    ["Evaluate/report", "evaluate_report.py → comparisons, sensitivity, calculation trace", "Same CSV; four- and five-criterion analyses"]], size=8.0)
p("The CLI does not reread the OSM files or the field observation template: those inputs are upstream of the generated CSV. The model receives the five criterion triples; identifiers and provenance metadata do not become extra TOPSIS criteria.")

doc.add_page_break()
h("V Implementation and Evaluation")
p("scripts/build_synthetic_dataset.py generates the synthetic CSV from the preserved OSM extract; the existing CLI/model pipeline evaluates it. scripts/evaluate_report.py regenerates reports/evaluation/summary.json, seeded weight runs, uncertainty scenarios, and a per-criterion trace. This builder reads that summary, including the input SHA-256 and runtime versions; numerical results are not hardcoded. Tests assess reproducibility and arithmetic consistency, not empirical validity.")
p(f"The recorded execution uses Python {evaluation['python']}, NumPy {evaluation['numpy']} and pandas {evaluation['pandas']}. Validation checks required metadata, unique IDs, coordinates, integer central scores and ordered finite bounds in 1–5; missing values are rejected rather than set to zero. The calculation is deterministic: randomness is confined to the documented synthetic generator and explicitly seeded sensitivity experiments. Appendix B reconstructs one complete five-criterion result.")
h("VI Results")
top = evaluation["top_candidates"]
table(["Rank", "OSM node ID", "Mapped context (hypothetical)", "Score", "4-criterion rank", "Crisp rank"], [
    [r["rank"], r["alternative_id"].removeprefix("SYN-OSM-"), r["location_name"].split(" — hypothetical")[0], f(r["score"]), r["base_rank"], r["crisp_rank"]] for r in top[:8]])
p("Table 1 reports the leading eight of 28 alternatives under the five-criterion model. Candidate IDs are SYN-OSM- followed by the displayed node ID; full IDs appear in Appendix A. The score and ranks describe this synthetic candidate set only; shared-bound normalization means adding or removing alternatives can change results. The comparisons below use the same alternatives.")
table(["Comparison", "Top-1 changed", "Top-3 overlap", "Spearman"], [[
    x["comparison"].replace("_", " "), x["top1_changed"], f(x["top_k_overlap"], 3), f(x["spearman"], 3)] for x in evaluation["comparisons"]])
p("The leading two candidates exchange their demand and vehicle fuzzy triples while their other three triples match. Equal weights yield equal scores; weight changes expose that trade-off rather than establishing a unique best location. Adding E1 exchanges the third and fourth baseline positions, so two of the three selected leading alternatives remain shared. This illustrates a controlled extension effect, not a before/after safety improvement. Rounded scores should not be used to infer meaningful separation of near-tied candidates.")
s = evaluation["sensitivity"]
doc.add_page_break()
h("Sensitivity to assumptions")
p("Weight sensitivity uses 1,000 seeded perturbations: each equal weight is multiplied by an independent U(0.8,1.2) factor and renormalized. These are exploratory scenarios, not stakeholder preference probabilities. Width sensitivity contracts triangles toward their centers by α∈{0,0.5,1} and recomputes normalization; it tests the full representation/scaling procedure, not fuzziness in isolation. Exact ties share minimum rank; top-choice frequencies use deterministic ID ordering, not probabilities.")
p(f"The five-criterion top group is tied between {top[0]['alternative_id']} and {top[1]['alternative_id']} at {f(top[0]['score'],4)}. Adding E1 gives top-three overlap {f(evaluation['comparisons'][1]['top_k_overlap'],3)} versus the four-criterion baseline (Spearman {f(evaluation['comparisons'][1]['spearman'],3)}); this is an assumption comparison, not validation.")
p(f"Across {evaluation['draws']} weight draws, mean top-three overlap with equal weights is {f(s['mean_top_k_overlap'],3)} (minimum {f(s['min_top_k_overlap'],3)}); mean Spearman rank correlation is {f(s['mean_spearman'],3)} (minimum {f(s['min_spearman'],3)}). These perturbations are scenarios, not probabilities of field priority.")
u = evaluation["uncertainty"]
p("Uncertainty-width scenarios (α, top-ranked ID(s), Spearman versus full width): " + "; ".join(
    f"{x['width_factor']}: {', '.join(x['top1_ids'])}, {f(x['spearman_vs_full'],3)}" for x in u) + ".")
FIGS.mkdir(parents=True, exist_ok=True)
fig, ax = plt.subplots(figsize=(6.4, 1.35))
freq = sorted((r for r in s["top1_frequencies"] if r["top1_count"]),
              key=lambda x: (-x["top1_fraction"], x["alternative_id"]))
bars = ax.barh([r["alternative_id"] for r in freq], [100*r["top1_fraction"] for r in freq], color="#315f89")
ax.invert_yaxis(); ax.set_xlim(0, 110); ax.set_xlabel("Weight draws selected as top-1 (%)", fontsize=9)
ax.tick_params(labelsize=9)
for bar, row in zip(bars, freq):
    ax.text(bar.get_width()+1, bar.get_y()+bar.get_height()/2, f"{100*row['top1_fraction']:.1f}%", va="center", fontsize=9)
fig.tight_layout(pad=.3)
fig.savefig(FIGS / "weight_top1_sensitivity.png", dpi=200); plt.close(fig)
doc.add_picture(str(FIGS / "weight_top1_sensitivity.png"), width=Inches(5.7))
p(f"Figure 1. Only the {len(freq)} alternatives selected at least once are shown; the other {evaluation['candidate_count']-len(freq)} candidates had zero selections. Deterministic tie handling applies. These are experiment frequencies, not probabilities of need.")

h("VII Discussion and Limitations")
p("The experiment demonstrates computational behavior, not relative safety. A high closeness score means proximity to the modeled high-priority ideal within this candidate set. The public map extract supports candidate geometry and contextual proxies, but 27 junction candidates and the one mapped crossing remain unverified. In particular, road class is not traffic volume or operating speed; mapped tags do not certify current conditions; unknown attributes remain unknown. The destination proxy is not a verified route or access measure, and shared amenity inputs may count context twice.")
p("Central-score demand/destination association is reported as a diagnostic, not evidence that criteria are empirically independent. The five-versus-four comparison tests the mathematical consequence of adding E1 at equal weight, not predictive accuracy or improved decisions. The crisp comparator changes representation and scaling. Sensitivity ranges are analyst-chosen, fuzzy bounds uncalibrated, and equal weights provisional. A stable or unstable rank under these scenarios cannot validate the assumptions.")
p("Before operational use, collect matched pedestrian and vehicle counts, independently documented accessibility/control audits, current photographs, walking-route and destination verification, dates and disruptions, costs, jurisdiction, and feasibility. Preserve missingness and disagreements, then obtain independent review against predefined evidence. The model cannot establish code compliance, crash risk, treatment effect, or a budget-optimal portfolio. A low rank does not imply safety; a high rank does not quantify benefit.")
h("VIII Conclusion")
p("A reproducible five-criterion simulation demonstrates how a declared fuzzy TOPSIS adaptation can organize hypothetical crossing-improvement candidates and expose sensitivity to weighting and range assumptions. Its evidence-informed geography is not field evidence: public map facts, derived indicators, and modeled scores are kept distinct. Results must not be presented as validated site priorities. The report retains an older four-alternative example separately for teaching and provides reproducible artifacts for later replacement with verified observations.")

doc.add_page_break()
h("References")
with (ROOT / "docs/source_register.csv").open(encoding="utf-8") as fh:
    sources = {r["id"]: r for r in csv.DictReader(fh)}
for n, key in enumerate(("S2", "S3", "S4", "S1", "S8", "S9", "S10", "S7", "S11"), 1):
    r = sources[key]; p(f"[{n}] {r['full_reference']} {r['url']}")
h("Appendix A. Complete synthetic five-criterion ranking")
all_candidates = evaluation["all_candidates"]
ranking_table = table(["Rank", "Alternative ID", "Score", "4-criterion rank", "Crisp rank"], [[r["rank"], r["alternative_id"], f(r["score"],6), r["base_rank"], r["crisp_rank"]] for r in all_candidates])
ranking_table.autofit = False
for column, width in zip(ranking_table.columns, (0.5, 2.15, 1.15, 1.4, 1.4)):
    column.width = Inches(width)
for row in ranking_table.rows:
    for cell, width in zip(row.cells, (0.5, 2.15, 1.15, 1.4, 1.4)):
        cell.width = Inches(width)
h("Appendix B. Worked five-criterion calculation")
w = evaluation["worked_example"]
p(f"Alternative {w['alternative_id']}: raw and normalized triples, equal weight, weighted triple, and criterion-level vertex distances are shown below. Distances sum across criteria; C = D−/(D+ + D−).")
table(["Criterion", "Raw (l,c,u)", "Normalized", "Weight", "Weighted", "d+", "d−"], [[
    c["criterion"], str([round(float(v),4) for v in c["raw"]]), str([round(float(v),4) for v in c["normalized"]]),
    f(c["weight"],3), str([round(float(v),4) for v in c["weighted"]]), f(c["distance_positive"],6), f(c["distance_negative"],6)
] for c in w["criteria"]])
p(f"Totals: D+={f(w['distance_positive'],8)}, D−={f(w['distance_negative'],8)}, C={f(w['score'],8)}.")
p("For the demand criterion, shared bounds are 1 and 5: (3,4,5) becomes (0.5,0.75,1), then weight 0.2 gives (0.1,0.15,0.2). Its positive-ideal distance is √[((0.1−0.2)²+(0.15−0.2)²+(0.2−0.2)²)/3] = 0.06454972. Repeating for each criterion and summing gives the totals above; 0.76370815/(0.35295551+0.76370815) = 0.68391959. These calculations explain a synthetic result, not measured safety.")
h("Appendix C. Separate four-alternative teaching illustration")
p("The original hypothetical_demo.csv is a separate small teaching illustration; its invented rows are not part of the 28-candidate map-informed experiment. Its rankings are regenerated by the original demo pipeline and kept distinct.")
table(["ID", "4-criterion score", "4-criterion rank", "Crisp rank", "5-criterion score", "5-criterion rank"], [[
    r["alternative_id"], f(r["base_score"],6), r["base_rank"], r["crisp_rank"],
    f(r["extension_score"],6), r["extension_rank"]
] for r in evaluation["old_demo"]])
p("The four criteria are pedestrian demand, vehicle exposure, accessibility deficiency, and control/visibility deficiency; the five-criterion extension adds essential-destination access. Unlike the main 1–5 dataset, the original example uses invented hourly rates and 0–3 audit ratings, normalized within its own four rows. Raw triples remain in data/hypothetical_demo.csv. HYP-C moves from fourth to first after E1 is added, illustrating rank reversal; this does not validate the larger experiment. Scores from the two candidate sets are not directly comparable.")
h("Appendix D. Reproduction")
p("From the repository root, generate the map-informed synthetic CSV with python scripts/build_synthetic_dataset.py; run the synthetic-input CLI using the documented command in README.md and its dedicated synthetic output directory; then run python scripts/build_report.py. The report builder invokes scripts/evaluate_report.py and writes reports/evaluation/summary.json plus reports/report_figures/weight_top1_sensitivity.png. The four-row demo remains separately available through python src/run_project.py --demo. Do not use either synthetic dataset as field input. See reports/synthetic_data_methodology.md and data/source_inventory.csv for exact inputs, provenance, formulas, and limitations.")

footer = sec.footer.paragraphs[0]
footer.alignment = 2
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
doc.core_properties.title = title
doc.core_properties.author = ""
doc.save(ROOT / "P20_Fuzzy_TOPSIS_Report.docx")
