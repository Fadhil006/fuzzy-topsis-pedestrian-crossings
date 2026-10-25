"""Build the revised paper from verified project inputs; preserve the original."""
from __future__ import annotations

import csv
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from evaluate_revised_report import build_revised_evaluation

TITLE = "A Fuzzy TOPSIS Framework for Prioritising Pedestrian Crossing Safety and Accessibility Improvements"
TEAM_NUMBER = "034"
AUTHORS = (("Mohammed Ashar Husain", "2024BCS0137"),
           ("Muhamed Fadhil", "2024BCS0093"),
           ("Anjo P S", "2024BCS0097"))
REPOSITORY_URL = "https://github.com/Fadhil006/fuzzy-topsis-pedestrian-crossings"
OUTPUT = ROOT / "P20_Fuzzy_TOPSIS_Report_Revised.docx"
FIGURES = ROOT / "reports/revised_figures"


def clean(text):
    return str(text).replace("\u2014", ": ").replace("\u2013", "-").replace("\u2011", "-")


def fmt(value, places=3):
    return f"{float(value):.{places}f}"


def build_report():
    result = build_revised_evaluation()
    e = result["primary"]
    additional = result["additional_baselines"]
    crisp = {r["alternative_id"]: r for r in additional["crisp_topsis"]}
    deficiency = {r["alternative_id"]: r for r in additional["deficiency_rule"]}
    comparison = {r["comparison"]: r for r in e["comparisons"] + additional["comparisons"]}
    s = e["sensitivity"]
    with (ROOT / "data/processed/pedestrian_crossings_fuzzy_topsis.csv").open(encoding="utf-8", newline="") as stream:
        inputs = {r["alternative_id"]: r for r in csv.DictReader(stream)}

    FIGURES.mkdir(parents=True, exist_ok=True)
    frequency = sorted((r for r in s["top1_frequencies"] if r["top1_count"]),
                       key=lambda r: (-r["top1_count"], r["alternative_id"]))
    figure, axis = plt.subplots(figsize=(6.1, 1.35))
    bars = axis.barh([r["alternative_id"].removeprefix("SYN-OSM-") for r in frequency],
                     [100 * r["top1_fraction"] for r in frequency], color="#315f89")
    axis.invert_yaxis()
    axis.set_xlim(0, 60)
    axis.set_xlabel("Weight scenarios selecting candidate first (%)", fontsize=9)
    axis.tick_params(labelsize=9)
    for bar, row in zip(bars, frequency):
        axis.text(bar.get_width() + .6, bar.get_y() + bar.get_height() / 2,
                  f"{100 * row['top1_fraction']:.1f}%", va="center", fontsize=9)
    figure.tight_layout(pad=.3)
    figure.savefig(FIGURES / "weight_top1_sensitivity.png", dpi=220)
    plt.close(figure)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.60)
    section.left_margin = section.right_margin = Inches(.65)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.02
    for name, size in (("Title", 17), ("Heading 1", 12), ("Heading 2", 10.5)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(3)
    # Remove theme decorations so headings and equations have predictable layout.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)

    def p(text, size=None):
        paragraph = doc.add_paragraph(clean(text))
        if size:
            for run in paragraph.runs:
                run.font.size = Pt(size)
        return paragraph

    def h(text, level=1):
        return doc.add_heading(clean(text), level=level)

    def caption(text):
        paragraph = p(text, 9)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.bold = True

    def table(headers, rows, widths=None, size=9):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Shading Accent 1"
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = clean(text)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            for cell, value in zip(t.add_row().cells, values):
                cell.text = clean(value)
        if widths:
            t.autofit = False
            for column, width in zip(t.columns, widths):
                column.width = Inches(width)
            for row in t.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Inches(width)
        for row in t.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1
                    for run in paragraph.runs:
                        run.font.name, run.font.size = "Times New Roman", Pt(size)
        return t

    def math_run(text):
        run = OxmlElement("m:r")
        node = OxmlElement("m:t")
        node.text = clean(text)
        run.append(node)
        return run

    def math_parts(parent, parts):
        for part in parts if isinstance(parts, (list, tuple)) else [parts]:
            parent.append(math_run(part) if isinstance(part, str) else part)

    def indexed(base, subscript=None, superscript=None):
        tag = "m:sSubSup" if subscript and superscript else ("m:sSub" if subscript else "m:sSup")
        node = OxmlElement(tag)
        expression = OxmlElement("m:e")
        math_parts(expression, base)
        node.append(expression)
        for label, value in (("m:sub", subscript), ("m:sup", superscript)):
            if value:
                child = OxmlElement(label)
                child.append(math_run(value))
                node.append(child)
        return node

    def fraction(numerator, denominator):
        node = OxmlElement("m:f")
        for tag, text in (("m:num", numerator), ("m:den", denominator)):
            child = OxmlElement(tag)
            math_parts(child, text)
            node.append(child)
        return node

    def root_fraction(numerator, denominator):
        radical = OxmlElement("m:rad")
        properties = OxmlElement("m:radPr")
        hide = OxmlElement("m:degHide")
        hide.set(qn("m:val"), "1")
        properties.append(hide)
        radical.append(properties)
        radical.append(OxmlElement("m:deg"))
        expression = OxmlElement("m:e")
        expression.append(fraction(numerator, denominator))
        radical.append(expression)
        return radical

    def equation(number, *parts):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.8), WD_TAB_ALIGNMENT.RIGHT)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run("\t")
        math = OxmlElement("m:oMath")
        for part in parts:
            math.append(math_run(part) if isinstance(part, str) else part)
        paragraph._p.append(math)
        paragraph.add_run(f"\t({number})")

    def page():
        doc.add_page_break()

    title = doc.add_paragraph(TITLE, "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(f"Team {TEAM_NUMBER}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name, roll_number in AUTHORS:
        p(f"{name} ({roll_number})").alignment = WD_ALIGN_PARAGRAPH.CENTER
    p("Study area: IIIT Kottayam, Kerala, India").alignment = WD_ALIGN_PARAGRAPH.CENTER
    h("Abstract")
    p(f"This study examines how destination context affects pedestrian crossing intervention priorities near IIIT Kottayam. An archived OpenStreetMap extract was used to identify {e['candidate_count']} candidate locations and derive road and destination indicators. Documented rules then generated synthetic ratings for pedestrian demand, vehicle exposure, accessibility deficiency, control and visibility deficiency, and essential destination access. A triangular fuzzy TOPSIS model was compared with its four-criterion version and simple reference rankings. The two leading candidates tie at {fmt(e['all_candidates'][0]['score'],4)}. Adding destination access changes one member of the leading three, with rank correlation {fmt(comparison['base_vs_extension']['spearman'])}. The leading three remain unchanged across 1,000 weight perturbations. These findings demonstrate a reproducible prioritisation experiment and identify influential modelling choices; they are not a field assessment of crossing safety.")
    p("Keywords: pedestrian crossings; accessibility; fuzzy TOPSIS; public map data; synthetic modelling")
    h("I Introduction")
    p("Pedestrian improvements must respond to both local barriers and the journeys a crossing may serve. A location near a school or bus stop may have different access consequences from an otherwise similar location on a lightly connected road. Comparing these concerns requires an explicit decision rule rather than a single severity score. Here, each alternative represents a generic crossing improvement package attached to a mapped candidate point. The intended users are campus facilities staff and the relevant road authority, who would assess locations and engineering feasibility before considering construction.")
    p("The research question is whether adding an essential destination proximity criterion changes priorities when the original demand, road context and deficiency criteria are retained. The contribution is a traceable experiment linking public map evidence to reproducible synthetic inputs, not a new TOPSIS variant or a claim of improved safety. The earlier four-alternative experiment is retained separately in Appendix D.")
    h("II Related Work")
    p("Related transport studies motivate the decision framing (Table 1). Their alternatives and evidence differ from the present candidate packages, so their findings are not used as local calibration. Chen [4] supplies the fuzzy TOPSIS lineage, while Jefmański and Sagan [5] show why conversion scales deserve scrutiny. The full text of the Indian midblock study [7] illustrates the observations needed for empirical safety analysis. Other comparisons use verified records and available abstracts where full text was unavailable; no unreported validation results are inferred.")
    caption("Table 1. Related transport studies and their relevance")
    table(["Study", "Problem, evidence and method", "Evaluation and transfer limit"], [
        ["Battistini et al. [1]", "Ten unsignalized crossings on two Bologna roads; inspection checklists and intervention decision support.", "Abstract reports use in road safety planning. The checklist is not an Indian compliance standard."],
        ["Majumdar et al. [2]", "Pedestrian satisfaction and TOPSIS prioritisation of sidewalk and crosswalk attributes.", "Prioritises attributes rather than packages at individual sites; no local data are transferred."],
        ["Alver and Katanalp [3]", "Gap acceptance at semicontrolled midblock crossings; integrated AHP and fuzzy logic.", "Behavioural modelling, not validation of this study's synthetic deficiency ratings."],
        ["Golakiya et al. [7]", "Video observations at four Indian midblock sites; time to collision as a surrogate safety measure.", "Evidence from observed interactions, not crash-outcome validation or local crossing-package priorities."],
    ], widths=(1.25, 2.85, 2.85), size=9)
    page()

    h("III Study Scope and Data")
    h("A. Candidate locations and public evidence", 2)
    p("The OSM snapshot [8], archived and verified on 24 October 2026, contains 431 elements. Processing retained 286 motor-road ways and 91 destination objects after limited name and proximity deduplication. Candidate selection used a 3 km radius from the mapped campus centre (9.7554364 N, 76.6495350 E), prioritising mapped crossings, shared-node road junctions and destination access candidates, with at least 80 m spacing. The selected set comprises 27 junction candidates and one mapped crossing. These are not 28 verified pedestrian crossings, and inclusion denotes simulation eligibility rather than engineering approval.")
    p("Mapped geometry and tags form the observed public layer; distances and destination intensities are derived indicators; criterion ratings and fuzzy ranges are synthetic. For candidate 716872373 on Pala - Valavoor - Uzhavoor Road, mapped destinations include St Marys Church at about 60 m, Alphonsa English Medium School Valavoor at 83 m, Valavoor School at 277 m and a post office at 286 m. These are great-circle distances to feature points, not verified walking routes. NATPAC [6] and government portal resources [9] provide background only: district totals were not converted into crossing counts or score calibration.")
    h("B. Criteria and their interpretation", 2)
    p("All five criteria are dimensionless integer ratings from 1 to 5, with larger values indicating greater modelled priority. TOPSIS treats successive categories as equally spaced. Table 2 summarises their meaning; Appendix E gives the complete generation rules. Costs are excluded because no verified estimates exist.")
    caption("Table 2. Increasing-priority criteria and their evidence basis")
    table(["Code and criterion", "Basis of the synthetic rating", "Reason for inclusion"], [
        ["P1 Pedestrian demand", "Mapped destination intensity, bus proximity, road context and seeded variation.", "Represents plausible crossing activity without inventing pedestrian counts."],
        ["P2 Vehicle exposure", "Adjacent road class and explicit mapped lane and speed-limit tags.", "Represents traffic context, not measured volume or operating speed."],
        ["P3 Accessibility deficiency", "Neutral prior with explicit wheelchair, kerb and tactile-tag adjustments.", "Represents potential accessibility barriers under incomplete evidence."],
        ["P4 Control/visibility deficiency", "Separate prior with explicit signals, markings and lighting adjustments.", "Represents control and visibility concerns separately from accessibility."],
        ["E1 Essential destination access", "Strongest nearby destination contribution and total proximity intensity.", "Adds access importance beyond aggregate destination intensity."],
    ], widths=(1.65, 2.65, 2.65), size=9)
    h("C. Preprocessing and fuzzy representation", 2)
    p("The generator preserves source references and distinguishes missing tags from explicit negative tags. The CLI checks required metadata, unique identifiers, coordinates, ordered finite triples and agreement between integer scores and their central values. Missing criterion values are rejected, not imputed as zero; excluded candidates require a recorded reason. All 28 generated records passed. Geometry-derived road nodes lack guaranteed child-node tags in this snapshot, and a wider refresh failed. Consequently, 27 candidates in each deficiency criterion rely on an assumption-based prior.")
    p("Each centre c is represented by (l,c,u), where l=max(1,c-h) and u=min(5,c+h). Width h is 1 for low-confidence proxies and 1.5 for very-low-confidence assumptions or incomplete context. These are deliberately wider modelling ranges for weaker evidence, not measured variation or calibrated confidence intervals. Triangular membership rises to one at c and falls to zero at the bounds; membership is not probability. The effect of these choices is tested through width contraction.")
    page()

    h("IV Methodology")
    h("A. Normalisation and weighting", 2)
    p("Let i index alternatives, j criteria and q the lower, central or upper component. Shared bounds Lⱼ=minᵢlᵢⱼ and Uⱼ=maxᵢuᵢⱼ place every triangle of a criterion on the same scale:")
    equation(1, indexed("z", "ijq"), " = ",
             fraction([indexed("x", "ijq"), " - ", indexed("L", "j")],
                      [indexed("U", "j"), " - ", indexed("L", "j")]), ",   q ∈ {l,c,u}.")
    p("Zero-range criteria are flagged, and criteria whose normalised triples do not distinguish alternatives are removed. Remaining weights are renormalised over n active criteria. All criteria are already priority-oriented, so no cost reversal is needed. Equal weights provide a transparent baseline without implying expert elicitation (Table 3).")
    caption("Table 3. Provisional criterion weights")
    table(["Model", "P1", "P2", "P3", "P4", "E1", "Sum"], [
        ["Four criteria", "0.25", "0.25", "0.25", "0.25", "Excluded", "1.00"],
        ["Five criteria", "0.20", "0.20", "0.20", "0.20", "0.20", "1.00"],
    ], size=9)
    equation(2, indexed("w", "j"), " = ", fraction("1", "n"), ",   ",
             indexed("w", "j"), " ≥ 0,   ", indexed("∑", "j"), indexed("w", "j"), " = 1,   ",
             indexed("ṽ", "ij"), " = ", indexed("w", "j"), indexed("z̃", "ij"), ".")
    h("B. Fuzzy TOPSIS distances and ranking", 2)
    p("For each active criterion, the positive ideal is Aⱼ⁺=(wⱼ,wⱼ,wⱼ) and the negative ideal Aⱼ⁻=(0,0,0). The positive ideal denotes greatest intervention priority, not the safest condition. The vertex distance compares all three triangle coordinates:")
    differences = []
    for component in ("l", "c", "u"):
        if differences:
            differences.append(" + ")
        differences.append(indexed(["(", indexed("a", component), " - ", indexed("b", component), ")"], superscript="2"))
    equation(3, "d(ã,b̃) = ", root_fraction(differences, "3"), ".")
    p("Criterion distances are summed, rather than combined as a single Euclidean distance across criteria:")
    equation(4, indexed("D", "i", "+"), " = ", indexed("∑", "j"), "d(", indexed("ṽ", "ij"), ",", indexed("A", "j", "+"), "),   ",
             indexed("D", "i", "-"), " = ", indexed("∑", "j"), "d(", indexed("ṽ", "ij"), ",", indexed("A", "j", "-"), ").")
    equation(5, indexed("C", "i"), " = ",
             fraction([indexed("D", "i", "-")], [indexed("D", "i", "+"), " + ", indexed("D", "i", "-")]), ".")
    p("Higher C indicates closer agreement with the high-priority profile. A zero denominator is rejected. Exact ties share the minimum rank, while IDs provide stable display order; near-equal floating-point values are ranked before rounding. The model uses Chen-style vertex distances and closeness [4], with explicitly adopted common-bound normalisation, crisp weights and weighted theoretical ideals. Appendix B reconstructs a complete result.")
    h("C. Destination extension and reference comparisons", 2)
    p("The extension adds E1 while leaving the first four inputs unchanged and preserving their relative weights. E1 emphasises the strongest nearby essential destination; P1 emphasises total intensity and bus proximity. The Spearman rank correlation between their central scores is " + fmt(e["ped_destination_spearman"], 6) + ", so the extension may partly count the same context twice. It is a proximity proxy, not an assessment of necessary walking connections.")
    p("The existing simple baseline averages separately normalised central scores for P1 to P4. Two additional references address the guidance's comparison design: an equal mean of normalised P3 and P4 centres (the deficiency-only rule), and conventional crisp TOPSIS on all five centres. Crisp TOPSIS uses observed weighted ideals and Euclidean distance across criteria, unlike the fuzzy model's summed vertex distances. Appendix C gives its worked calculation; comparisons do not isolate uncertainty alone.")
    page()

    h("V Implementation and Evaluation")
    h("A. Reproducible processing", 2)
    p(f"Two stages separate data construction from ranking. The generator reads the archived OSM response and study metadata, then writes a collated CSV. The CLI reads only that CSV and assembles a 28 × 5 × 3 decision array; identifiers, coordinates and provenance are not extra criteria. The recorded environment is Python {e['python']}, NumPy {e['numpy']} and pandas {e['pandas']}. Generation uses seed {e['data_seed']}, with a separate random stream for each OSM node; weight sensitivity uses seed {e['sensitivity_seed']}. Input hashes, cleaned data, processing logs and criterion traces make the outputs auditable. The project repository [10] records the implementation and supporting files; Appendix F identifies the files and commands.")
    h("B. Comparison and validation design", 2)
    p("Top-three overlap counts shared members divided by three; Spearman correlation compares average ranks. Exact-tie boundaries use deterministic ID order for selection statistics. Independent arithmetic checks, input tests and repeat runs assess implementation consistency. No held-out site records, independent condition audit or expert review are available; model agreement is therefore not empirical validation.")
    h("VI Results")
    h("A. Equal-weight priorities and extension effect", 2)
    caption("Table 4. Leading candidates under five-criterion fuzzy TOPSIS")
    table(["Rank", "OSM node ID", "Fuzzy score", "Four-criterion rank", "Mean rule rank", "Crisp TOPSIS rank"], [
        [r["rank"], r["alternative_id"].removeprefix("SYN-OSM-"), fmt(r["score"], 4), r["base_rank"], r["crisp_rank"], crisp[r["alternative_id"]]["rank"]]
        for r in e["top_candidates"][:5]
    ], widths=(.55, 1.55, 1.05, 1.25, 1.2, 1.35), size=9)
    p("Full candidate IDs begin SYN-OSM-; Appendix A lists all 28. The leading pair, 716872373 and 764885861 on Pala - Valavoor - Uzhavoor Road, exchange their demand and vehicle triples while the other three triples match. Their equal score is therefore explained by equal weighting, rather than evidence that the sites have identical conditions. Adding E1 exchanges the third and fourth positions of the four-criterion model, replacing one member of the leading three.")
    caption("Table 5. Ranking comparisons on the same 28 candidates")
    labels = [("Four-criterion fuzzy vs mean rule", "base_vs_crisp"),
              ("Four-criterion vs five-criterion fuzzy", "base_vs_extension"),
              ("Five-criterion fuzzy vs deficiency rule", "five_criterion_fuzzy_vs_deficiency_rule"),
              ("Five-criterion fuzzy vs crisp TOPSIS", "five_criterion_fuzzy_vs_crisp_topsis")]
    table(["Comparison", "Top choice changed", "Top-three overlap", "Spearman ρ"], [
        [label, "Yes" if comparison[key]["top1_changed"] else "No", fmt(comparison[key]["top_k_overlap"]), fmt(comparison[key]["spearman"])]
        for label, key in labels
    ], widths=(2.9, 1.4, 1.4, 1.25), size=9)
    p("The crisp TOPSIS reference retains the fuzzy model's leading three, whereas the deficiency-only rule shares none. The latter favours accessibility and control ratings without demand or road exposure. This contrast shows that the decision framing matters more than the small changes produced by adding E1 in this experiment. The original four-criterion mean remains a distinct comparator; it is not crisp TOPSIS.")
    page()

    h("B. Sensitivity to weights and fuzzy widths", 2)
    p("For each of 1,000 weight scenarios, independent factors drawn uniformly from [0.8,1.2] multiply the equal weights, which are then renormalised. The leading three remain unchanged: mean and minimum overlap are both " + fmt(s["mean_top_k_overlap"]) + ". Mean rank correlation is " + fmt(s["mean_spearman"]) + ", with minimum " + fmt(s["min_spearman"]) + ". The two equal-weight leaders divide the first-place selections between " + str(frequency[0]["top1_count"]) + " and " + str(frequency[1]["top1_count"]) + " runs (Figure 1). Changing the relative emphasis on demand and vehicle exposure resolves their initial trade-off. Exact remaining ties use ID order for the first-place count.")
    doc.add_picture(str(FIGURES / "weight_top1_sensitivity.png"), width=Inches(5.7))
    caption("Figure 1. First-place selections in the five-criterion weight experiment. Labels are OSM node IDs; the other 26 candidates received zero selections. Frequencies describe these scenarios, not probabilities of safety or need.")
    p("Width sensitivity moves both bounds towards the same centre and refits normalisation:")
    equation(6, indexed("x̃", "ij"), "(α) = ", indexed("c", "ij"), " + α(", indexed("x̃", "ij"), " - ", indexed("c", "ij"), "),   α ∈ {0,0.5,1}.")
    caption("Table 6. Five-criterion width scenarios relative to the full-width ranking")
    table(["Width factor α", "Interpretation", "Spearman ρ", "Leading pair"], [
        [fmt(r["width_factor"], 1), {0.0: "Centres only", .5: "Half widths", 1.0: "Original widths"}[r["width_factor"]], fmt(r["spearman_vs_full"], 4), "Unchanged"]
        for r in e["uncertainty"]
    ], widths=(1.25, 2.3, 1.3, 2.1), size=9)
    p("The leading pair persists in all three scenarios. Lower positions change modestly, reflecting both altered widths and refitted bounds. This is a test of the full representation and scaling procedure, not evidence that the chosen widths are calibrated.")
    h("VII Discussion and Limitations")
    p("The results support a transparent way to organise candidate review and identify measurements that could change priorities. Stable leading groups under the selected perturbations coexist with much weaker agreement with the deficiency-only rule. Future assessment should therefore examine both road activity and local barriers, rather than interpret a stable numerical rank as sufficient evidence for action.")
    p("The principal limitation is the assumed rating model. Mapping completeness, candidate selection, equal ordinal spacing, fuzzy widths and equal weights all affect the result. In particular, missing node tags leave most deficiency ratings prior-based, and E1 overlaps strongly with demand context. Source timestamps do not establish current conditions or usable routes. Before operational use, verify crossing existence, ownership and feasible packages; collect matched pedestrian and vehicle counts, dated condition audits and route evidence; and obtain independent review. Costs and dependencies would also be needed for budget allocation. No crash reduction, legal compliance or intervention benefit is estimated.")
    h("VIII Conclusion")
    p("A reproducible public-map-informed simulation evaluated 28 hypothetical pedestrian crossing improvement candidates near IIIT Kottayam. Adding destination proximity changed one member of the leading three without separating the tied leaders. The leading group was stable under the tested weight and width changes, but differed substantially from a deficiency-only priority rule. The contribution is an inspectable link from mapped context through synthetic assumptions to ranked results. Independent observations are still needed before those results can support site-specific engineering decisions.")
    h("Acknowledgement")
    p("The study uses data from OpenStreetMap contributors under ODbL 1.0 and the public research and government resources cited below.", 9)
    page()

    h("References")
    references = [
        '[1] R. Battistini, C. Lantieri, A. Simone, G. Dondi, and V. Vignali, "A Decision Support System for the safety evaluation of urban pedestrian crossings," Transportation Research Procedia, vol. 69, pp. 655-662, 2023, doi: 10.1016/j.trpro.2023.02.220.',
        '[2] B. B. Majumdar, P. K. Sahu, M. Patil, and N. Vendotti, "Pedestrian Satisfaction-Based Methodology for Prioritization of Critical Sidewalk and Crosswalk Attributes Influencing Walkability," Journal of Urban Planning and Development, vol. 147, no. 3, art. 04021032, 2021, doi: 10.1061/(ASCE)UP.1943-5444.0000718.',
        '[3] Y. Alver and B. Y. Katanalp, "Evaluating the Pedestrian Gap Acceptance in Semicontrolled Midblock Crosswalks with an Integrated AHP-FL Approach," Journal of Transportation Engineering, Part A: Systems, vol. 148, no. 9, art. 04022066, 2022, doi: 10.1061/JTEPBS.0000720.',
        '[4] C.-T. Chen, "Extensions of the TOPSIS for group decision-making under fuzzy environment," Fuzzy Sets and Systems, vol. 114, no. 1, pp. 1-9, 2000, doi: 10.1016/S0165-0114(97)00377-1.',
        '[5] B. Jefmański and A. Sagan, "Item Response Theory Models for the Fuzzy TOPSIS in the Analysis of Survey Data," Symmetry, vol. 13, no. 2, art. 223, 2021, doi: 10.3390/sym13020223.',
        '[6] NATPAC, Draft Road Safety Vision: Kerala, historical policy report, undated. Accessed Oct. 24, 2026. https://natpac.kerala.gov.in/wp-content/uploads/2023/08/Draft-Road-Safety-Vision.pdf.',
        '[7] H. Golakiya, R. Chauhan, C. S. Bari, and A. Dhamaniya, "Pedestrian safety analysis at urban midblock section under mixed traffic conditions using time to collision as surrogate safety measure," Current Science, vol. 123, no. 9, pp. 1117-1128, 2022, doi: 10.18520/cs/v123/i9/1117-1128.',
        '[8] OpenStreetMap contributors, "OpenStreetMap campus context and Nominatim campus lookup," snapshot archived Oct. 24, 2026. https://www.openstreetmap.org/copyright. Data licence: ODbL 1.0. Exact queries and source metadata are included in the project archive.',
        '[9] Kerala Data Portal, "District-wise PWD road-length and motor-vehicle chart resources," Government of Kerala. Accessed Oct. 24, 2026. https://datahub.kerala.gov.in/datasets?page=12.',
        f'[10] Team {TEAM_NUMBER}, "Fuzzy TOPSIS pedestrian crossings: implementation, data and reports," GitHub repository. [Online]. Available: {REPOSITORY_URL}.',
    ]
    for reference in references:
        p(reference, 9.5)

    h("Appendix A. Complete main-experiment rankings")
    p("The following values preserve the original 28-candidate results and add two explicitly computed reference rankings. The mean rule uses P1 to P4 central values; crisp TOPSIS uses all five. Exact ties and floating-point near ties are assigned before rounding.")
    caption("Table A1. Full ranking and correctly distinguished reference ranks")
    table(["Fuzzy rank", "Alternative ID", "Fuzzy score", "Fuzzy rank (4)", "Mean rank (4)", "Crisp rank (5)", "Deficiency rank"], [
        [r["rank"], r["alternative_id"], fmt(r["score"], 6), r["base_rank"], r["crisp_rank"], crisp[r["alternative_id"]]["rank"], deficiency[r["alternative_id"]]["rank"]]
        for r in e["all_candidates"]
    ], widths=(.5, 2.1, .95, .8, .8, .9, .9), size=8.5)
    page()

    h("Appendix B. Worked fuzzy TOPSIS calculation")
    w = e["worked_example"]
    p(f"For {w['alternative_id']}, all five shared fuzzy outer bounds in this dataset are 1 and 5. Demand (3,4,5) therefore becomes (0.5,0.75,1); weight 0.20 produces (0.1,0.15,0.2). Each criterion has ideal (0.2,0.2,0.2) and opposite ideal (0,0,0). Table B1 shows the complete calculation.")
    caption("Table B1. Fuzzy matrix row, weighted values and criterion distances")
    table(["Criterion", "Raw (l,c,u)", "Normalised", "Weight", "Weighted", "d+", "d-"], [
        [c["criterion"], str([round(v, 3) for v in c["raw"]]), str([round(v, 3) for v in c["normalized"]]), fmt(c["weight"], 2), str([round(v, 3) for v in c["weighted"]]), fmt(c["distance_positive"], 6), fmt(c["distance_negative"], 6)]
        for c in w["criteria"]
    ], widths=(.8, 1.25, 1.25, .6, 1.25, .9, .9), size=8.5)
    p("For demand, d+ = √[((0.1-0.2)²+(0.15-0.2)²+(0.2-0.2)²)/3] = 0.06454972. Its negative-ideal distance is √[(0.1²+0.15²+0.2²)/3] = 0.15545632.")
    p(f"Summing all five rows gives D+={fmt(w['distance_positive'],8)} and D-={fmt(w['distance_negative'],8)}. Hence C=0.76370815/(0.35295551+0.76370815)={fmt(w['score'],8)}, matching the primary output. The independent two-criterion arithmetic test in tests/test_model.py separately checks the implemented vertex-distance summation.")
    h("Appendix C. Worked conventional crisp TOPSIS reference")
    c = result["worked_crisp"]
    p("The same candidate has centres (4,5,3,4,4). Each central column is normalised using its own observed minimum and maximum across the 28 candidates, not the fuzzy outer bounds. Apply equal active weights and define A+ as the column maxima of the weighted central matrix, with A- as its minima. Distances are D+=√[Σⱼ(vᵢⱼ-Aⱼ+)²] and D-=√[Σⱼ(vᵢⱼ-Aⱼ-)²]. This standard crisp procedure is distinct from both fuzzy vertex summation and the original central-score mean rule.")
    caption("Table C1. Complete crisp calculation for SYN-OSM-716872373")
    table(["Code", "Centre", "Min", "Max", "Normalised", "Weight", "Weighted", "Ideal +", "Ideal -"], [
        [name, fmt(c["rawcentral"][j], 1), fmt(c["minimums"][j], 1), fmt(c["maximums"][j], 1), fmt(c["normalized"][j], 4), fmt(c["weights"][j], 2), fmt(c["weighted"][j], 4), fmt(c["positiveideal"][j], 2), fmt(c["negativeideal"][j], 2)]
        for j, name in enumerate(("P1", "P2", "P3", "P4", "E1"))
    ], size=8.5)
    p(f"The weighted row is (0.15,0.20,0.13333333,0.20,0.20). Its squared shortfall from the positive ideal is 0.05²+0.06666667². Thus D+={fmt(c['distance_positive'],8)}, D-={fmt(c['distance_negative'],8)} and C={fmt(c['score'],8)}. The different closeness scale does not imply that the crisp model is more accurate.")
    page()

    h("Appendix D. Supplementary four-alternative experiment")
    p("The earlier hypothetical_demo.csv experiment is a separate reproducibility example. Its invented hourly count ranges and 0 to 3 ratings differ from the main 1 to 5 scale, and the two datasets are never pooled. The mean-rule ranks below refer to the original four-criterion central-score mean, not conventional crisp TOPSIS.")
    caption("Table D1. Preserved supplementary experiment results")
    table(["ID", "Four-fuzzy score", "Four-fuzzy rank", "Four-mean rank", "Five-fuzzy score", "Five-fuzzy rank"], [
        [r["alternative_id"], fmt(r["base_score"], 6), r["base_rank"], r["crisp_rank"], fmt(r["extension_score"], 6), r["extension_rank"]]
        for r in e["old_demo"]
    ], size=9)
    p("HYP-C moves from fourth to first after E1 is added, whereas the larger experiment retains its two leading candidates. The examples demonstrate different responses to the same extension; closeness scores from different candidate sets are not directly comparable.")
    h("Appendix E. Exact synthetic generation rules")
    p("For a destination at great-circle distance d≤500 m, let its contribution be a(1-d/600). The destination weight a is 3 for hospitals, clinics, schools and kindergartens; 2.5 for colleges, universities, pharmacies, bus stations and marketplaces; 2 for bus stops or public transport platforms/stop positions; 1.5 for banks, post offices, community centres and places of worship; and 1 for other shops, restaurants and cafes. I sums contributions, M is their maximum (zero when none occur), and B=1 if a bus stop/platform/stop position is within 250 m. Same-named features within 100 m are deduplicated; aliases can remain. Area features use approximate geometry centres rather than entrances.")
    p("Let R(x)=min(5,max(1,floor(x+0.5))). Per-candidate streams use random.Random(f\"20261024:{osm_node_id}\"). The highest-hierarchy adjacent road is chosen, breaking ties by lower way ID. Its base b is 5 for motorway/trunk/primary, 4 secondary, 3 tertiary, 2 unclassified/residential and 1 living street/service (including link variants). P2=R(b+0.5·1[lanes≥3]+0.5·1[maxspeed≥60]+U(-0.35,0.35)); adjustments require explicit numeric tags. P1=R(1+1.05ln(1+I)+0.35B+0.12(P2-2)+U(-0.65,0.65)). E1=R(1+1.05M+0.25ln(1+I)). These coefficients are declared modelling choices, not estimated effects.")
    p("Draw shared s~U(-0.45,0.45), then separate normal errors with mean 0 and standard deviation 0.85. P3=R(3+s+error+access adjustments). Wheelchair yes/designated subtracts 1.2; explicit no adds 1.2. Flush/lowered kerb or sloped_curb=yes subtracts 0.8; raised kerb adds 0.8. Tactile paving yes subtracts 0.4; explicit no adds 0.4. P4=R(3+s+0.2(P2-3)+separate error+control adjustments). Candidate-node traffic signals subtract 1.1; otherwise explicit zebra/marked crossing tags subtract 0.7. Mapped lighting at the node or selected road subtracts 0.4. Missing tags make no adjustment. The source script fixes draw order and preserves the tags that triggered each adjustment.")
    p("For each rating centre, h=1 unless the criterion's confidence label begins very_low, in which case h=1.5; clip l=c-h and u=c+h to [1,5]. The wider category applies to deficiency criteria without recognised tag evidence, E1 with no mapped destination within 500 m, and both P1 and E1 when the query buffer is truncated or unknown. These are declared evidence-quality categories, not statistical confidence levels. All selected 500 m disks lie inside the saved query rectangle, but map completeness is unknown. The supplementary experiment retains its original independently defined triples.")
    page()

    h("Appendix F. Reproduction and supporting package")
    repository = p("Project repository: ")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), doc.part.relate_to(REPOSITORY_URL, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    link_style = OxmlElement("w:rStyle")
    link_style.set(qn("w:val"), "Hyperlink")
    properties.append(link_style)
    run.append(properties)
    link_text = OxmlElement("w:t")
    link_text.text = REPOSITORY_URL
    run.append(link_text)
    hyperlink.append(run)
    repository._p.append(hyperlink)
    p("Use Python 3.10 or later and the recorded project dependencies in a local virtual environment. Authoring additionally uses python-docx and matplotlib. Source refresh is optional; the saved inputs allow offline regeneration. The generator reads data/raw/osm_context.json and data/raw/osm_metadata.json. The ranking CLI reads data/processed/pedestrian_crossings_fuzzy_topsis.csv only.")
    for command in (
        "python -m venv .venv",
        "source .venv/bin/activate",
        "python -m pip install -r requirements.txt",
        "python -m pip install python-docx matplotlib",
        "python scripts/build_synthetic_dataset.py",
        "python src/run_project.py --synthetic-input data/processed/pedestrian_crossings_fuzzy_topsis.csv --output-dir outputs",
        "python scripts/build_revised_report.py",
        "python -m unittest discover -s tests -v",
    ):
        p(command, 9)
    caption("Table F1. Supporting files and their roles")
    table(["Artifact", "Location and purpose"], [
        ["Raw sources and metadata", "data/raw/; exact successful and failed queries, public resource bodies and metadata. The OSM base timestamp is not the HTTP completion time."],
        ["Source records", "data/source_inventory.csv and docs/source_register.csv; URLs, archive/access dates, coverage, licence notes and bibliographic records."],
        ["Main inputs", "data/processed/pedestrian_crossings_fuzzy_topsis.csv and synthetic_generation_manifest.json; scores, intervals, provenance, assumptions and hashes."],
        ["Cleaned data and log", "outputs/synthetic_cleaned_decision_matrix.csv and synthetic_preprocessing_log.csv; included/excluded status and retained notes."],
        ["Primary ranking", "outputs/fuzzy_topsis_results.csv; all five criteria. Original outputs/synthetic_* diagnostics retain their documented four-criterion scope."],
        ["Revised evaluation", "reports/revised_evaluation/; original five-criterion diagnostics in primary/, new reference rankings, comparison CSV and summary.json."],
        ["Mathematics and tests", "src/model.py, scripts/evaluate_report.py, scripts/evaluate_revised_report.py, tests/ and docs/worked_example.md."],
        ["Report and extended method", "P20_Fuzzy_TOPSIS_Report_Revised.docx; scripts/build_revised_report.py; reports/synthetic_data_methodology.md."],
    ], widths=(1.6, 5.35), size=9)
    p("The saved OSM response uses bounds 9.7284, 76.6225, 9.7825 and 76.6765 (south, west, north, east). The attempted wider extract and member-node refresh failed, so the original snapshot was retained. The actual observation CSV remains empty. The public sources and source-derived coordinates must not be represented as crossing surveys; future field input uses a separate CLI mode.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = TITLE
    doc.core_properties.author = "; ".join(name for name, _ in AUTHORS)
    doc.core_properties.subject = "Pedestrian crossing prioritisation near IIIT Kottayam"
    for part in doc.part.package.parts:
        if hasattr(part, "element"):
            for text in part.element.iter():
                if text.text and any(dash in text.text for dash in ("\u2013", "\u2014")):
                    raise ValueError("Forbidden punctuation in generated document")
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build_report()
